
import openpyxl
from collections import defaultdict, Counter

EXCEL_PATH = r"C:\Users\Finstein-Emp\Documents\projects\Dataverse_v4\requirements-excel-file\dataverse-Testcases-V4.xlsx"

def main():
    sep = "=" * 100
    dash = "-" * 80
    hash_line = "#" * 100

    print(sep)
    print("LOADING:", EXCEL_PATH)
    print(sep)

    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

    print()
    print(sep)
    print("SHEET NAMES (%d sheets)" % len(wb.sheetnames))
    print(sep)
    for i, name in enumerate(wb.sheetnames, 1):
        print("  %d. %s" % (i, name))

    global_urs_utc_map = defaultdict(set)
    global_utc_count = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print()
        print(hash_line)
        print("SHEET: '%s'" % sheet_name)
        print(hash_line)

        all_rows = []
        for row in ws.iter_rows(values_only=True):
            all_rows.append(tuple(row))

        if not all_rows:
            print("  (EMPTY SHEET)")
            continue

        header_row_idx = None
        headers = []
        for idx, row in enumerate(all_rows):
            non_none = [c for c in row if c is not None]
            if len(non_none) >= 2:
                header_row_idx = idx
                headers = list(row)
                break

        if header_row_idx is None:
            print("  (NO HEADER ROW FOUND)")
            for idx, row in enumerate(all_rows[:5]):
                print("  Row %d: %s" % (idx, str(row)))
            continue

        headers_clean = []
        for h in headers:
            if h is not None:
                headers_clean.append(str(h).strip())
            else:
                headers_clean.append(None)

        data_rows = all_rows[header_row_idx + 1:]
        data_rows = [r for r in data_rows if any(c is not None for c in r)]

        print()
        print("  TOTAL DATA ROWS: %d" % len(data_rows))
        print("  HEADER ROW INDEX: %d" % header_row_idx)
        print("  TOTAL COLUMNS: %d" % len(headers_clean))

        print()
        print("  COLUMN HEADERS:")
        print("  " + dash)
        for ci, h in enumerate(headers_clean):
            if h is not None:
                print("    Col %d: '%s'" % (ci, h))

        key_keywords = ["urs", "srs", "sds", "utc", "test case", "title", "description",
                        "given", "when", "then", "expected", "test type", "priority",
                        "pre-condition", "precondition", "step", "result", "status",
                        "module", "feature", "category", "requirement", "scenario"]

        print()
        print("  KEY COLUMNS DETECTED:")
        print("  " + dash)
        for ci, h in enumerate(headers_clean):
            if h is None:
                continue
            h_lower = h.lower()
            for kw in key_keywords:
                if kw in h_lower:
                    print("    Col %d: '%s' (matched: '%s')" % (ci, h, kw))
                    break

        print()
        print("  FIRST 5 DATA ROWS:")
        print("  " + dash)
        for ri, row in enumerate(data_rows[:5]):
            print()
            print("  --- Row %d ---" % (ri + 1))
            for ci, val in enumerate(row):
                if val is not None and ci < len(headers_clean):
                    col_name = headers_clean[ci] if headers_clean[ci] else "(unnamed col %d)" % ci
                    val_str = str(val)
                    if len(val_str) > 200:
                        val_str = val_str[:200] + "..."
                    print("    %s: %s" % (col_name, val_str))

        urs_col = None
        srs_col = None
        sds_col = None
        utc_col = None

        for ci, h in enumerate(headers_clean):
            if h is None:
                continue
            h_lower = h.lower().strip()
            if "urs" in h_lower and urs_col is None:
                urs_col = ci
            if "srs" in h_lower and srs_col is None:
                srs_col = ci
            if "sds" in h_lower and sds_col is None:
                sds_col = ci
            if "utc" in h_lower and utc_col is None:
                utc_col = ci

        if urs_col is not None or utc_col is not None:
            print()
            print("  URS-UTC MAPPING (URS col=%s, SRS col=%s, SDS col=%s, UTC col=%s):" % (urs_col, srs_col, sds_col, utc_col))
            print("  " + dash)

            urs_utc_map = defaultdict(set)
            urs_srs_map = defaultdict(set)
            utc_values = set()

            for row in data_rows:
                urs_val = None
                srs_val = None
                utc_val = None

                if urs_col is not None and urs_col < len(row) and row[urs_col] is not None:
                    urs_val = str(row[urs_col]).strip()
                if srs_col is not None and srs_col < len(row) and row[srs_col] is not None:
                    srs_val = str(row[srs_col]).strip()
                if utc_col is not None and utc_col < len(row) and row[utc_col] is not None:
                    utc_val = str(row[utc_col]).strip()

                if urs_val and utc_val:
                    urs_utc_map[urs_val].add(utc_val)
                    global_urs_utc_map[urs_val].add(utc_val)
                if urs_val and srs_val:
                    urs_srs_map[urs_val].add(srs_val)
                if utc_val:
                    utc_values.add(utc_val)
                    global_utc_count += 1

            if urs_utc_map:
                for urs in sorted(urs_utc_map.keys()):
                    utcs = sorted(urs_utc_map[urs])
                    srs_list = sorted(urs_srs_map.get(urs, set()))
                    srs_display = str(srs_list[:5])
                    if len(srs_list) > 5:
                        srs_display += "..."
                    print("    %s: %d UTCs, SRS=%s" % (urs, len(utcs), srs_display))
                    for u in utcs[:3]:
                        print("      - %s" % u)
                    if len(utcs) > 3:
                        print("      ... and %d more" % (len(utcs) - 3))
            else:
                print("    Total unique UTCs in this sheet: %d" % len(utc_values))

        categorical_keywords = ["test type", "priority", "status", "category", "module"]
        for ci, h in enumerate(headers_clean):
            if h is None:
                continue
            h_lower = h.lower()
            for kw in categorical_keywords:
                if kw in h_lower:
                    vals = Counter()
                    for row in data_rows:
                        if ci < len(row) and row[ci] is not None:
                            vals[str(row[ci]).strip()] += 1
                    if vals:
                        print()
                        print("  UNIQUE VALUES for '%s' (%d unique):" % (h, len(vals)))
                        for v, count in vals.most_common(20):
                            print("    '%s': %d" % (v, count))
                    break

    print()
    print(sep)
    print("GLOBAL SUMMARY")
    print(sep)
    print()
    print("Total sheets: %d" % len(wb.sheetnames))
    print("Total UTC entries (across all sheets): %d" % global_utc_count)
    print("Total unique URS IDs: %d" % len(global_urs_utc_map))
    print()
    print("URS -> UTC COUNT (ALL SHEETS COMBINED):")
    print("-" * 60)
    total_unique_utcs = 0
    for urs in sorted(global_urs_utc_map.keys()):
        utcs = global_urs_utc_map[urs]
        total_unique_utcs += len(utcs)
        print("  %s: %d unique UTCs" % (urs, len(utcs)))
    print()
    print("TOTAL UNIQUE UTCs: %d" % total_unique_utcs)

    wb.close()
    print()
    print("Done.")

if __name__ == "__main__":
    main()
